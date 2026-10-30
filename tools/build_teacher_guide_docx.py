from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "代码说明_教师测试指南.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(9)
    if color:
        r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, bold=True, color=(255, 255, 255))
        shade(t.rows[0].cells[i], "1F4E78")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return t


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    for i, line in enumerate(text.strip("\n").splitlines()):
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        if i < len(text.strip("\n").splitlines()) - 1:
            r.add_break()
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [("Heading 1", 16, "1F4E78"), ("Heading 2", 13, "2F5597"), ("Heading 3", 11, "404040")]:
        styles[name].font.name = "Microsoft YaHei"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(color)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LLM 辅助 Web 漏洞挖掘与验证\n教师详细测试指南")
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(31, 78, 121)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p2.add_run("项目：llm-web-vuln-lab    版本：main    日期：2026-10-30")
    rr.font.size = Pt(10)
    rr.font.color.rgb = RGBColor(100, 100, 100)
    doc.add_paragraph()

    doc.add_heading("一 项目概述", level=1)
    doc.add_paragraph("本项目用于网络安全课程实验，演示如何使用大语言模型辅助进行 SQL 注入（SQLi）和跨站脚本（XSS）代码审计，并通过人工复核、本地教学靶场和受控 POC 验证模型结论。项目不扫描公网系统，所有动态验证仅面向本机 127.0.0.1:4280 的授权教学环境。")
    doc.add_paragraph("当前仓库 main 分支已包含 24 个教学样本、36 条正式模型记录、26 条加分实验记录、实验报告、验证证据和自动化检查工具。")

    doc.add_heading("二 小组成员与代码作者", level=1)
    table(doc, ["标识", "作者", "主要负责内容"], [
        ("A", "刘珂如", "漏洞判定标准、真值复核、指标统计、实验整合、最终报告"),
        ("B", "张思佳", "SQLi/XSS Prompt 设计、V0-V1-V2 迭代、模型对照实验"),
        ("C", "常佳佳", "样本构造、源码预处理、输入包、行号映射、数据追溯"),
        ("D", "李楠", "Docker/DVWA 环境、POC 验证、Burp 联动、运行证据"),
    ], [0.6, 1.1, 4.8])
    doc.add_paragraph("历史批量生成的样本、模型原文和证据文件不重复改写，以免改变哈希和实验可追溯性；这些文件通过作者索引和查询脚本归属作者。")

    doc.add_heading("三 运行错误时查询作者", level=1)
    doc.add_paragraph("老师如果运行过程中遇到错误，请先记录错误信息中的相对文件路径，再在仓库根目录执行下面的命令：")
    code(doc, "python tools/show_author.py <报错文件路径>")
    doc.add_paragraph("例如，若错误来自 environment/check_runtime.py：")
    code(doc, "python tools/show_author.py environment/check_runtime.py")
    doc.add_paragraph("脚本会输出文件路径、作者、负责内容和对应操作手册。主要目录对应关系如下：")
    table(doc, ["报错路径", "作者"], [
        ("tools/、tests/", "A 刘珂如"),
        ("prompts/", "B 张思佳"),
        ("dataset/、submission_inputs/", "C 常佳佳"),
        ("environment/、pocs/、evidence/、demo/", "D 李楠"),
        ("docs/A_*、results/A_*", "A 刘珂如"),
        ("docs/B_*、results/B_*", "B 张思佳"),
        ("docs/C_*、results/C_*", "C 常佳佳"),
        ("docs/D_*、results/D_*", "D 李楠"),
        ("runs/、experiments/、templates/", "全组协作"),
    ], [3.4, 3.1])
    doc.add_paragraph("完整路径映射位于根目录 AUTHORS.md；查询程序位于 tools/show_author.py。共同文件的最终整合责任由 A 刘珂如承担，具体运行责任以文件元数据和阶段前缀为准。")

    doc.add_heading("四 目录与文件说明", level=1)
    table(doc, ["目录", "用途", "主要作者"], [
        ("dataset/", "24 个 SQLi/XSS 教学样本、原始源码、真值和输入包", "C"),
        ("prompts/", "SQLi、XSS V0/V1/V2 及高级案例 Prompt", "B"),
        ("runs/", "模型输入、原始回答、规范化 JSON 和运行元数据", "全组"),
        ("pocs/", "模型生成的原始 POC、适配 POC 和复核脚本", "D"),
        ("evidence/", "请求、响应、截图、日志和人工复核证据", "D"),
        ("experiments/", "Prompt Injection、Burp、消融、二阶 SQLi、IDOR", "全组"),
        ("environment/", "Docker/DVWA、数据库初始化和运行检查", "D"),
        ("tools/", "冻结检查、提交审计、指标评估和证据索引", "A"),
        ("docs/", "实验报告、操作手册、判定标准和测试指南", "全组/A整合"),
    ], [1.3, 4.6, 0.8])

    doc.add_heading("五 测试环境准备", level=1)
    doc.add_paragraph("静态检查只需要 Git 和 Python 3.10 以上；动态验证还需要 Docker Desktop、Docker Compose 和浏览器。Windows PowerShell 使用 python，Linux/macOS/WSL 可将其替换为 python3。")
    code(doc, "python --version\ngit --version\ndocker version\ndocker compose version")
    doc.add_paragraph("推荐从 GitHub 的 main 分支获取代码：")
    code(doc, "git clone https://github.com/krrb2006/llm-web-vuln-lab.git\ncd llm-web-vuln-lab\ngit branch --show-current")
    doc.add_paragraph("若仅进行静态检查和单元测试，不需要安装 PHP，也不需要启动 Docker。")

    doc.add_heading("六 最小静态检查流程", level=1)
    doc.add_paragraph("请先在仓库根目录依次执行：")
    code(doc, "python -m unittest discover -s tests -v\npython tools/check_freeze.py --report-only\npython tools/audit_submission.py")
    table(doc, ["检查项", "正常结果", "作用"], [
        ("项目测试", "Ran 22 tests ... OK", "验证冻结、统计、输入防篡改和安全边界工具"),
        ("冻结检查", "结论：READY", "检查 24 个样本、36 条正式记录、Prompt 哈希和行号映射"),
        ("提交审计", "各项 READY", "检查正式实验、加分实验和最终报告证据是否齐全"),
    ], [1.3, 1.8, 3.6])
    doc.add_paragraph("这些命令只读仓库文件，不调用模型、不部署靶场、不执行 POC。如果遇到 Python 导入错误，请确认当前目录是仓库根目录，而不是仓库的上一级目录。")

    doc.add_heading("七 阅读实验结果", level=1)
    doc.add_paragraph("建议按以下顺序阅读：")
    for item in [
        "docs/最终实验报告_2026-10-26.md：最终结论、指标、限制和复现命令。",
        "docs/00_总执行手册.md：实验流程和成员职责。",
        "results/formal_eval_2026-10-19-v2/summary.md：正式 36 条记录的统计。",
        "results/evidence_index_2026-10-21.md：运行 ID 到证据文件的索引。",
    ]:
        bullet(doc, item)
    doc.add_paragraph("正式实验为 12 个测试样本分别使用 V0、V1、V2 Prompt，共 36 条记录；加分实验包括 16 条 Prompt Injection、2 条 Burp 联动、4 条上下文消融和 4 条高级案例。")

    doc.add_heading("八 可选启动本地 DVWA 教学靶场", level=1)
    doc.add_paragraph("动态验证只允许在本机授权教学环境执行。首次使用时准备官方 DVWA：")
    code(doc, "git clone https://github.com/digininja/DVWA.git third_party/DVWA\n$aRef = 'd498903da01d89a695fa49d5392d3776eb901e2f'\n$cRef = '8da63f7b8e169e5037626b0362d005ba7ba5ede0'\n$dRef = 'a296e4b527ebdaf41ca04c2bbc5627b1a89efc06'\npython environment/build_lab_cases.py --a-ref $aRef --c-ref $cRef --d-ref $dRef\npython environment/init_lab_db.py")
    doc.add_paragraph("启动服务并确认端口只绑定到 127.0.0.1:4280：")
    code(doc, "docker compose --env-file local/lab.env -f third_party/DVWA/compose.yml -f environment/compose.lab.yml up -d\ndocker compose --env-file local/lab.env -f third_party/DVWA/compose.yml -f environment/compose.lab.yml ps")
    doc.add_paragraph("浏览器访问 http://127.0.0.1:4280/。全新 DVWA 数据卷首次启动时，按页面提示完成 Setup DVWA，然后返回登录页。")
    doc.add_paragraph("检查服务和样本语法：")
    code(doc, "python environment/check_runtime.py --a-ref $aRef --c-ref $cRef --d-ref $dRef\npython environment/check_sql_http.py --sample S009 --sample S010\npython environment/check_xss_browser.py\npython environment/check_d_business.py")
    doc.add_paragraph("语法检查、HTTP 200 或脚本退出码 0 都不能单独证明漏洞成立；必须结合普通请求、受控请求、源码数据流和预先定义的观察判据。")

    doc.add_heading("九 单个样本的复核方法", level=1)
    for item in [
        "在 dataset/manifest.json 中找到样本编号、漏洞类型和真值记录。",
        "在 dataset/raw/<sample_id>/ 阅读原始 PHP、样本卡和真值复核说明。",
        "在 dataset/packets/ 查看发送给模型的固定输入和 line_map.json。",
        "在 runs/F_<sample>_<version>_R1/ 查看模型原文、规范化 JSON、输入哈希和运行元数据。",
        "在 pocs/ 和 evidence/ 查看 POC、请求响应、截图和本地验证记录。",
        "对照 results/ 中的复核记录，区分模型预测、代码定位、原始 POC 可执行和实际漏洞触发。",
    ]:
        bullet(doc, item)

    doc.add_heading("十 常见问题与处理", level=1)
    table(doc, ["现象", "原因", "处理方式"], [
        ("ModuleNotFoundError: tools...", "未在仓库根目录运行", "cd 到包含 tools、tests、dataset 的仓库根目录后重试"),
        ("Docker 无法连接", "Docker Desktop 未启动或 WSL 未集成", "启动 Docker Desktop，确认 docker version 正常"),
        ("端口 4280 被占用", "已有服务占用教学端口", "停止占用端口的本地服务后重试，不要改成公网地址"),
        ("数据库 Access denied", "DVWA 或教学数据库未初始化", "查看 environment/README.md，重新确认 lab.env 和 compose 配置"),
        ("PHP 语法错误", "样本文件或挂载目录不一致", "查询作者后联系对应成员，先核对源码 SHA-256，不要直接修改冻结样本"),
        ("作者不明确", "错误路径不在当前目录映射中", "执行 python tools/show_author.py <路径>，再查看 AUTHORS.md"),
    ], [1.6, 2.2, 3.1])

    doc.add_heading("十一 停止服务与提交注意事项", level=1)
    code(doc, "docker compose --env-file local/lab.env -f third_party/DVWA/compose.yml -f environment/compose.lab.yml down")
    for item in [
        "不要使用 docker compose down -v，以免删除 DVWA 数据卷。",
        "不要提交 local/lab.env、登录 Cookie、令牌或未脱敏截图。",
        "不要把模型回答、HTTP 200、脚本退出码 0 单独当作漏洞触发证据。",
        "正式统计只使用冻结后的 36 条记录；加分实验不回写主实验指标。",
        "项目结果只代表本组小规模、受控教学样本，不外推为真实世界漏洞发现率。",
    ]:
        bullet(doc, item)

    doc.add_heading("十二 交付文件清单", level=1)
    table(doc, ["文件", "用途"], [
        ("README.md", "项目入口和快速开始"),
        ("AUTHORS.md", "完整目录到作者映射"),
        ("tools/show_author.py", "根据错误文件路径查询作者"),
        ("docs/代码说明_教师测试指南.md", "本指南的 Markdown 源稿"),
        ("docs/最终实验报告_2026-10-26.md", "最终实验报告源稿"),
        ("docs/网络安全大作业实验报告.docx", "实验报告 Word 版"),
        ("docs/网络安全大作业实验报告.pdf", "实验报告 PDF 版"),
    ], [3.1, 3.8])
    doc.add_paragraph("本指南由 A 刘珂如负责最终整合，代码分工为 A 刘珂如、B 张思佳、C 常佳佳、D 李楠。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
